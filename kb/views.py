import csv
import io

from django.contrib import messages
from django.shortcuts import render, redirect

from django.utils import timezone

from openpyxl import load_workbook

from rest_framework.decorators import api_view
from rest_framework.response import Response

from .forms import CSVUploadForm
from .models import URLDocument

from .services.ingestion import run_ingestion_pipeline
from .services.scraper import scrape_url


# ============================================================
# HOME PAGE
# ============================================================

def home(request):

    return render(
        request,
        "kb/home.html"
    )


# ============================================================
# UPLOAD CSV / XLSX
# ============================================================

def upload_csv(request):

    if request.method == "POST":

        form = CSVUploadForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            uploaded_file = request.FILES["file"]

            file_name = uploaded_file.name.lower()

            urls = []

            # ------------------------------------------------
            # CSV
            # ------------------------------------------------

            if file_name.endswith(".csv"):

                decoded_file = (
                    uploaded_file
                    .read()
                    .decode("utf-8")
                )

                csv_file = io.StringIO(
                    decoded_file
                )

                reader = csv.DictReader(
                    csv_file
                )

                if not reader.fieldnames:

                    messages.error(
                        request,
                        "The CSV file is empty."
                    )

                    return redirect(
                        "upload_csv"
                    )

                # Find URL column
                url_column = None

                for column in reader.fieldnames:

                    if (
                        column.strip().lower()
                        == "url"
                    ):

                        url_column = column
                        break

                if not url_column:

                    messages.error(
                        request,
                        "The CSV must contain a 'URL' column."
                    )

                    return redirect(
                        "upload_csv"
                    )

                # Extract URLs
                for row in reader:

                    url = row.get(
                        url_column,
                        ""
                    ).strip()

                    if url:

                        urls.append(url)

            # ------------------------------------------------
            # XLSX
            # ------------------------------------------------

            elif file_name.endswith(".xlsx"):

                workbook = load_workbook(
                    uploaded_file,
                    read_only=True,
                    data_only=True
                )

                worksheet = workbook.active

                rows = worksheet.iter_rows(
                    values_only=True
                )

                try:

                    headers = next(rows)

                except StopIteration:

                    messages.error(
                        request,
                        "The Excel file is empty."
                    )

                    workbook.close()

                    return redirect(
                        "upload_csv"
                    )

                # Find URL column
                url_index = None

                for index, header in enumerate(
                    headers
                ):

                    if (
                        header is not None
                        and str(header)
                        .strip()
                        .lower()
                        == "url"
                    ):

                        url_index = index
                        break

                if url_index is None:

                    messages.error(
                        request,
                        "The Excel file must contain a 'URL' column."
                    )

                    workbook.close()

                    return redirect(
                        "upload_csv"
                    )

                # Extract URLs
                for row in rows:

                    if url_index >= len(row):

                        continue

                    value = row[url_index]

                    if value is not None:

                        url = str(
                            value
                        ).strip()

                        if url:

                            urls.append(url)

                workbook.close()

            # ------------------------------------------------
            # Save URLs to database
            # ------------------------------------------------

            added = 0

            for url in urls:

                _, created = (
                    URLDocument.objects
                    .get_or_create(
                        url=url
                    )
                )

                if created:

                    added += 1

            # ------------------------------------------------
            # AUTOMATIC INGESTION
            # ------------------------------------------------

            try:

                pipeline_result = (
                    run_ingestion_pipeline()
                )

                scrape_results = (
                    pipeline_result[
                        "scrape_results"
                    ]
                )

                total_chunks = (
                    pipeline_result[
                        "total_chunks"
                    ]
                )

                successful = sum(
                    1
                    for result in scrape_results
                    if result["success"]
                )

                failed = (
                    len(scrape_results)
                    - successful
                )

                messages.success(
                    request,
                    (
                        f"{added} new URLs added. "
                        f"{successful} URLs scraped successfully, "
                        f"{failed} failed. "
                        f"{total_chunks} chunks indexed."
                    )
                )

            except Exception as error:

                messages.error(
                    request,
                    f"Ingestion failed: {error}"
                )

            return redirect(
                "upload_csv"
            )

    else:

        form = CSVUploadForm()

    return render(
        request,
        "kb/upload.html",
        {
            "form": form
        }
    )


# ============================================================
# OLD MANUAL SCRAPE FUNCTION
# ============================================================
# Kept for now so we don't break anything.
# Our new ingestion pipeline is preferred.

def scrape_documents():

    documents = URLDocument.objects.all()

    results = []

    for document in documents:

        result = scrape_url(
            document.url
        )

        document.status_code = (
            result["status_code"]
        )

        document.title = (
            result["title"]
        )

        document.raw_html = (
            result["raw_html"]
        )

        document.clean_text = (
            result["clean_text"]
        )

        document.scraped_at = (
            timezone.now()
        )

        document.save()

        results.append({
            "url": document.url,
            "success": result["success"],
            "status_code": result["status_code"],
            "title": result["title"],
            "error": result["error"],
        })

    return results


# ============================================================
# SEARCH API
# ============================================================

@api_view(["GET"])
def search_api(request):

    question = request.GET.get(
        "q",
        ""
    ).strip()

    if not question:

        return Response(
            {
                "error":
                    "Please provide a search query using ?q="
            },
            status=400
        )

    try:

        from .services.rag import (
            answer_question
        )

        result = answer_question(
            question
        )

        return Response(
            result
        )

    except Exception as error:

        return Response(
            {
                "error":
                    f"Search failed: {str(error)}"
            },
            status=500
        )


# ============================================================
# URLS API
# ============================================================

@api_view(["GET"])
def urls_api(request):

    documents = (
        URLDocument.objects
        .all()
        .order_by("id")
    )

    data = []

    for document in documents:

        data.append({

            "url": document.url,

            "http_status":
                document.status_code,

            "title":
                document.title,

            "raw_html":
                document.raw_html,

            "clean_text":
                document.clean_text,

            "metadata": {

                "scraped_at":
                    document.scraped_at,

                "created_at":
                    document.created_at,

                "updated_at":
                    document.updated_at,
            }
        })

    return Response(
        data
    )